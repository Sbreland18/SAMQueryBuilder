"""Formatted Excel copies (.xlsx): bold frozen header row, filters on every column, sensible widths, IDs and
codes kept as text so leading zeros survive."""
from __future__ import annotations

import io
import re

_NUM = re.compile(r"^-?\d{1,15}(\.\d+)?$")


def to_xlsx(headings: list, rows: list, sheet: str = "Report", text_columns: set | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]\*\?/\\:]", " ", sheet)[:31] or "Report"
    ws.append([str(h) for h in headings])
    keep_text = set(text_columns or ())
    for i, h in enumerate(headings):  # IDs, codes, zips and phones stay text (leading zeros)
        if re.search(r"id$|code|zip|phone|number|grade|ssn|lunch", str(h), re.I):
            keep_text.add(i)
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    formula_like = []
    for r in rows:
        out = []
        for i, v in enumerate(r):
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)     # control characters Excel can't store
            if v is None:
                out.append(None)
            elif i not in keep_text and isinstance(v, str) and _NUM.match(v.strip()) and not v.startswith("0"):
                out.append(float(v) if "." in v else int(v))
            else:
                out.append(v)
        ws.append(out)
        row_no = ws.max_row
        for i, v in enumerate(out):     # text starting with = + - @ must stay text, never a formula
            if isinstance(v, str) and v[:1] in "=+-@":
                formula_like.append((row_no, i + 1, v))
    for rr, cc, v in formula_like:
        cell = ws.cell(row=rr, column=cc)
        cell.value = v
        cell.data_type = "s"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF2F7")
    for c in ws[1]:
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if headings:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headings))}{max(1, len(rows) + 1)}"
    for i, h in enumerate(headings, 1):
        sample = [len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None]
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, max(sample) + 2))
        if i - 1 in keep_text:
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_to_xlsx(csv_path, xlsx_path=None) -> str:
    import csv
    from pathlib import Path
    p = Path(csv_path)
    with open(p, newline="", encoding="utf-8-sig", errors="replace") as fh:
        rows = list(csv.reader(fh))
    out = Path(xlsx_path) if xlsx_path else p.with_suffix(".xlsx")
    out.write_bytes(to_xlsx(rows[0] if rows else [], rows[1:], p.stem))
    return str(out)
