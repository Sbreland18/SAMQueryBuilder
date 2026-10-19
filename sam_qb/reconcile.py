"""Compare with the vendor's side: match the file you send (from SAM) with the vendor's own export or error
report and list who is missing on either side - and, optionally, values that differ (e.g. emails)."""
from __future__ import annotations

import csv
import re
from pathlib import Path


def read(path) -> tuple[list[str], list[list[str]]]:
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(p, read_only=True, data_only=True).active
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = p.read_text(encoding="utf-8-sig", errors="replace")
        sample = text[:4000]
        delim = "\t" if sample.count("\t") > sample.count(",") else ("|" if sample.count("|") > sample.count(",")
                                                                    else ",")
        rows = list(csv.reader(text.splitlines(), delimiter=delim))
    rows = [r for r in rows if any(c.strip() for c in r)]
    return (rows[0], rows[1:]) if rows else ([], [])


def _norm(v) -> str:
    v = str(v or "").strip().lower()
    return re.sub(r"^0+(?=\d)", "", v) if v.isdigit() else v


def _col(rows, i):
    return {_norm(r[i]) for r in rows if i < len(r) and str(r[i]).strip()}


def suggest_keys(h1, r1, h2, r2) -> tuple[int, int, float]:
    """(our column, their column, share of our values found on their side) for the best-matching ID pair."""
    best = (0, 0, 0.0)
    idlike = re.compile(r"id|number|code|email|user|sourcedid|login", re.I)
    for i, a in enumerate(h1):
        va = _col(r1[:3000], i)
        if not va or (not idlike.search(a) and len(va) < 0.8 * max(1, len(r1[:3000]))):
            continue
        for j, b in enumerate(h2):
            vb = _col(r2[:3000], j)
            if not vb:
                continue
            share = len(va & vb) / len(va)
            bonus = 0.05 if idlike.search(b) else 0
            if share + bonus > best[2] + (0.05 if idlike.search(h2[best[1]] if h2 else "") else 0):
                best = (i, j, share)
    return best


def compare(h1, r1, h2, r2, key1: int, key2: int, field_pairs: list[tuple[int, int]] = ()) -> dict:
    """{'only_ours': [rows], 'only_theirs': [rows], 'matched': n, 'differences': [(key, ours heading,
        our value, their value)], 'dupes_ours': n, 'dupes_theirs': n}"""
    ours, theirs = {}, {}
    d1 = d2 = 0
    for r in r1:
        k = _norm(r[key1]) if key1 < len(r) else ""
        if k:
            d1 += k in ours
            ours.setdefault(k, r)
    for r in r2:
        k = _norm(r[key2]) if key2 < len(r) else ""
        if k:
            d2 += k in theirs
            theirs.setdefault(k, r)
    only_ours = [ours[k] for k in ours if k not in theirs]
    only_theirs = [theirs[k] for k in theirs if k not in ours]
    diffs = []
    for k in ours.keys() & theirs.keys():
        a, b = ours[k], theirs[k]
        for i, j in field_pairs:
            va = a[i] if i < len(a) else ""
            vb = b[j] if j < len(b) else ""
            if _norm(va) != _norm(vb):
                diffs.append((a[key1], h1[i], va, vb))
    return {"only_ours": only_ours, "only_theirs": only_theirs, "matched": len(ours.keys() & theirs.keys()),
            "differences": sorted(diffs), "dupes_ours": d1, "dupes_theirs": d2,
            "total_ours": len(ours), "total_theirs": len(theirs)}


def save_report(path, h1, h2, res: dict):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["Where", "Detail"] + list(h1))
        for r in res["only_ours"]:
            w.writerow(["In SAM file only", "missing at the vendor"] + list(r))
        w.writerow([])
        w.writerow(["Where", "Detail"] + list(h2))
        for r in res["only_theirs"]:
            w.writerow(["At the vendor only", "not in the SAM file"] + list(r))
        w.writerow([])
        w.writerow(["Different values", "ID", "Field", "SAM file", "Vendor"])
        for k, f, a, b in res["differences"]:
            w.writerow(["Different", k, f, a, b])
